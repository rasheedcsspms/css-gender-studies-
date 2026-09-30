"""Low-level Word helpers: document setup, styles, inline markup, boxes, cards, tables, fields."""
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Inches

from . import style as S

CONTENT_WIDTH_CM = 16.6  # A4 (21 cm) minus 2 x 2.2 cm margins


def _rgb(h):
    return RGBColor.from_string(h)


# ============================================================ document setup
def new_document(accent=S.PLUM):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)
    _styles(doc, accent)
    _settings(doc)
    return doc


def _font(st, name=S.BODY_FONT, size=None, bold=None, italic=None, colour=None):
    f = st.font
    f.name = name
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if colour:
        f.color.rgb = _rgb(colour)


def _get_style(doc, name, base="Normal", kind=1):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, kind)
        st.base_style = doc.styles[base]
        return st


def _border_bottom(pPr, colour, sz=8, space=4):
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", str(space)), ("w:color", colour)):
        b.set(qn(k), v)
    pbdr.append(b)
    pPr.append(pbdr)


def _styles(doc, accent):
    st = doc.styles["Normal"]
    _font(st, size=S.BODY_SIZE, colour=S.INK)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15

    heads = {
        1: (20, accent, 0, 12),
        2: (15, S.TEAL, 16, 6),
        3: (12.5, S.NAVY, 12, 4),
        4: (11.5, S.CORAL, 10, 3),
        5: (11, S.SLATE, 8, 2),
    }
    for lvl, (size, col, before, after) in heads.items():
        h = doc.styles[f"Heading {lvl}"]
        _font(h, S.HEAD_FONT, size, True, lvl == 5, col)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.line_spacing = 1.05
        if lvl == 1:
            h.paragraph_format.page_break_before = True
            _border_bottom(h.element.get_or_add_pPr(), accent, 12, 6)

    for name, indent in (("List Bullet", 0.63), ("List Bullet 2", 1.26), ("List Bullet 3", 1.89)):
        b = doc.styles[name]
        _font(b, size=S.BODY_SIZE, colour=S.INK)
        b.paragraph_format.space_after = Pt(3)
        b.paragraph_format.line_spacing = 1.12

    n = _get_style(doc, "GS Numbered")
    n.paragraph_format.left_indent = Cm(0.95)
    n.paragraph_format.first_line_indent = Cm(-0.95)
    n.paragraph_format.space_after = Pt(4)
    n.paragraph_format.tab_stops.add_tab_stop(Cm(0.95))

    lead = _get_style(doc, "GS Lead")
    _font(lead, size=12, italic=True, colour=S.SLATE)
    lead.paragraph_format.space_after = Pt(10)

    cap = doc.styles["Caption"]
    _font(cap, size=9, italic=True, bold=False, colour=S.SLATE)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(12)

    tc = _get_style(doc, "GS Table Caption", "Caption")
    tc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tc.paragraph_format.space_before = Pt(8)
    tc.paragraph_format.space_after = Pt(3)
    tc.paragraph_format.keep_with_next = True

    bx = _get_style(doc, "GS Box Text")
    _font(bx, size=10, colour=S.INK)
    bx.paragraph_format.space_after = Pt(3)
    bx.paragraph_format.line_spacing = 1.1

    bl = _get_style(doc, "GS Box Bullet", "GS Box Text")
    bl.paragraph_format.left_indent = Cm(0.5)
    bl.paragraph_format.first_line_indent = Cm(-0.35)

    tb = _get_style(doc, "GS Table Text")
    _font(tb, size=9.5, colour=S.INK)
    tb.paragraph_format.space_after = Pt(1)
    tb.paragraph_format.line_spacing = 1.05

    fig = _get_style(doc, "GS Figure")
    fig.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fig.paragraph_format.space_before = Pt(6)
    fig.paragraph_format.space_after = Pt(0)
    fig.paragraph_format.keep_with_next = True

    for lvl in (1, 2, 3):
        t = _get_style(doc, f"toc {lvl}")
        _font(t, size=10.5 if lvl > 1 else 11, bold=(lvl == 1), colour=S.INK)
        t.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
        t.paragraph_format.space_after = Pt(2 if lvl > 1 else 4)


def _settings(doc):
    """Ask Word to refresh fields (contents, page numbers) when the file is opened."""
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def set_properties(doc, title, subject, author="CSS Gender Studies notes"):
    cp = doc.core_properties
    cp.title, cp.subject, cp.author = title, subject, author
    cp.comments = subject
    cp.keywords = "CSS; Gender Studies; FPSC"


# ============================================================ inline markup
# **term** = key term (bold, subject colour) · ==date/number== (bold navy)
# ^^bold^^ plain bold · *italic*
_INLINE = re.compile(r"\*\*(.+?)\*\*|==(.+?)==|\^\^(.+?)\^\^|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])")


def add_runs(p, text, size=None, colour=None, bold=False, italic=False, accent=S.GREEN):
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            _run(p, text[pos:m.start()], size, colour, bold, italic)
        term, num, strong, ital = m.groups()
        if term is not None:
            _run(p, term, size, accent, True, italic)
        elif num is not None:
            _run(p, num, size, S.NAVY, True, italic)
        elif strong is not None:
            _run(p, strong, size, colour, True, italic)
        else:
            _run(p, ital, size, colour, bold, True)
        pos = m.end()
    if pos < len(text):
        _run(p, text[pos:], size, colour, bold, italic)
    return p


def _run(p, text, size=None, colour=None, bold=False, italic=False):
    r = p.add_run(text)
    if size:
        r.font.size = Pt(size)
    if colour:
        r.font.color.rgb = _rgb(colour)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    return r


def plain(text):
    """Strip inline markup (for word counts and captions)."""
    return _INLINE.sub(lambda m: next(g for g in m.groups() if g is not None), text)


# ============================================================ cell helpers
def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_borders(cell, top=None, left=None, bottom=None, right=None):
    """Each edge: (size_in_eighths_pt, 'RRGGBB') or None for no border."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        if spec:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec[0]))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec[1])
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)
    tcPr.append(borders)


def cell_margins(cell, top=80, bottom=80, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def set_col_widths(table, widths_cm):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_cm):
            gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, c in enumerate(row.cells):
            if i < len(widths_cm):
                c.width = Cm(widths_cm[i])


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trPr.append(el)


def table_indent_zero(table):
    tblPr = table._tbl.tblPr
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tblPr.append(ind)


def spacer(doc, pts=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run()
    r.font.size = Pt(pts)
    p.paragraph_format.line_spacing = Pt(pts)
    return p


# ============================================================ fields
def add_field(p, instr, placeholder=""):
    def fc(t):
        r = p.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), t)
        r._r.append(el)
        return r
    fc("begin")
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    if placeholder:
        p.add_run(placeholder)
    fc("end")


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def header_footer(doc, left_text, right_text, accent):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    hp = sec.header.paragraphs[0]
    from docx.shared import Twips
    hp.paragraph_format.tab_stops.add_tab_stop(Twips(4680), WD_TAB_ALIGNMENT.CLEAR)
    hp.paragraph_format.tab_stops.add_tab_stop(Twips(9360), WD_TAB_ALIGNMENT.CLEAR)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(CONTENT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    _run(hp, left_text, 8.5, S.SLATE)
    _run(hp, "\t" + right_text, 8.5, accent, bold=True)
    _border_bottom(hp._p.get_or_add_pPr(), S.tint(accent, 0.6), 4, 4)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = _run(fp, "", 9, S.SLATE)
    add_field(fp, "PAGE", "1")
    for run in fp.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = _rgb(S.SLATE)
