"""GSM — the Gender Studies Markup — parsed straight into a Word document.

See gs_build/SPEC.md §3 for the full syntax. In short:
  # .. #####        headings (H1 starts a new page)
  plain lines       a paragraph (consecutive lines are joined)
  - / "  - "        bullets (three levels)
  1. text           numbered paragraph (number kept as written)
  > text            lead paragraph (italic, larger)
  ::: box | title   a box   (simple example trap eye debate pakistan remember balance islam note glance update method)
  ::: card head     a card  (define thinker quote data law case report)  head = "Name | descriptor"
  ::: table | cap   a table (rows split on " | "; first row = header; @widths 3,7)
  ::: fig kind | c  a figure from a template (see figures.py)
  :::               closes a block
  @toc  @pagebreak  @py NAME args  @include file  %% comment
Inline: **key term**  ==date/number==  ^^bold^^  *italic*
"""
import re
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, Inches

from . import docxkit as K
from . import figures as F
from . import style as S


def read_source(path: Path):
    """Read a .gsm file, resolving @include recursively. Returns (front_matter, lines)."""
    lines = path.read_text(encoding="utf-8").splitlines()
    fm = {}
    if lines and lines[0].strip() == "---":
        end = lines.index("---", 1)
        for ln in lines[1:end]:
            k, _, v = ln.partition(":")
            fm[k.strip()] = v.strip()
        lines = lines[end + 1:]
    out = []
    for ln in lines:
        if ln.startswith("@include "):
            _, sub = read_source(path.parent / ln.split(None, 1)[1].strip())
            out.extend(sub)
        else:
            out.append(ln)
    return fm, out


class Builder:
    def __init__(self, doc, code, accent=S.PLUM, blocks=None):
        self.doc = doc
        self.code = code
        self.accent = accent
        self.fig_no = 0
        self.table_no = 0
        self.blocks = blocks or {}
        self.words = 0

    # ------------------------------------------------------------ counting
    def _count(self, text):
        self.words += len(K.plain(text).split())

    # ------------------------------------------------------------ primitives
    def heading(self, level, text):
        self._count(text)
        return self.doc.add_heading(text, level)

    def para(self, text, style=None, align=None):
        self._count(text)
        p = self.doc.add_paragraph(style=style)
        K.add_runs(p, text, accent=self.term_colour)
        if align == "centre":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    @property
    def term_colour(self):
        return S.GREEN

    def bullet(self, text, level=1):
        st = "List Bullet" if level == 1 else f"List Bullet {level}"
        return self.para(text, st)

    def numbered(self, num, text):
        self._count(text)
        p = self.doc.add_paragraph(style="GS Numbered")
        K._run(p, f"{num}\t", bold=True, colour=self.accent)
        K.add_runs(p, text, accent=self.term_colour)
        return p

    # ------------------------------------------------------------ boxes & cards
    def _cell(self, fill, edge_col, left_sz=28):
        t = self.doc.add_table(rows=1, cols=1)
        K.table_indent_zero(t)
        K.set_col_widths(t, [K.CONTENT_WIDTH_CM])
        c = t.cell(0, 0)
        K.shade(c, fill)
        K.cell_borders(c, top=(4, S.tint(edge_col, 0.55)), bottom=(4, S.tint(edge_col, 0.55)),
                       left=(left_sz, edge_col), right=(4, S.tint(edge_col, 0.55)))
        K.cell_margins(c, 100, 100, 180, 160)
        return t, c

    def _fill_cell(self, c, lines, col):
        para_buf = []

        def flush():
            if para_buf:
                txt = " ".join(para_buf)
                p = c.add_paragraph(style="GS Box Text")
                K.add_runs(p, txt, accent=self.term_colour)
                self._count(txt)
                para_buf.clear()

        for raw in lines:
            s = raw.strip()
            if not s:
                flush()
                continue
            if s.startswith("-> "):
                flush()
                p = c.add_paragraph(style="GS Box Text")
                p.paragraph_format.space_before = Pt(3)
                K._run(p, "➜  ", bold=True, colour=col)
                K.add_runs(p, s[3:], bold=True, colour=col, accent=col)
                self._count(s[3:])
            elif s.startswith("- "):
                flush()
                p = c.add_paragraph(style="GS Box Bullet")
                K._run(p, "•  ", bold=True, colour=col)
                K.add_runs(p, s[2:], accent=self.term_colour)
                self._count(s[2:])
            elif re.match(r"^\d+\.\s", s):
                flush()
                n, _, rest = s.partition(" ")
                p = c.add_paragraph(style="GS Box Bullet")
                K._run(p, n + "  ", bold=True, colour=col)
                K.add_runs(p, rest, accent=self.term_colour)
                self._count(rest)
            elif s.startswith('"') and s.endswith('"') and len(s) > 40:
                flush()
                p = c.add_paragraph(style="GS Box Text")
                K.add_runs(p, s, italic=True, size=10.5)
                self._count(s)
            else:
                para_buf.append(s)
        flush()
        # remove the empty first paragraph Word cells start with
        first_p = c.paragraphs[0]
        if not first_p.text and len(c.paragraphs) > 1:
            first_p._p.getparent().remove(first_p._p)

    def box(self, kind, title, lines):
        label, col = S.BOXES[kind]
        t, c = self._cell(S.tint(col, 0.93), col)
        p = c.paragraphs[0]
        p.style = self.doc.styles["GS Box Text"]
        K._run(p, label, 9, col, bold=True)
        if title:
            K._run(p, "  ·  " + title.upper() if kind != "glance" else "  ·  " + title.upper(), 9, col, bold=True)
            self._count(title)
        p.paragraph_format.space_after = Pt(4)
        holder = c.add_paragraph()
        self._fill_cell(c, lines, col)
        holder._p.getparent().remove(holder._p)
        if sum(len(x) for x in lines) < 900:
            K.no_split(t.rows[0])
        K.spacer(self.doc, 6)

    def card(self, kind, head, lines):
        label, col = S.CARDS[kind]
        name, _, desc = head.partition("|")
        t, c = self._cell("FFFFFF", col, 36)
        K.cell_borders(c, top=(6, S.tint(col, 0.4)), bottom=(6, S.tint(col, 0.4)), left=(36, col),
                       right=(6, S.tint(col, 0.4)))
        p = c.paragraphs[0]
        p.style = self.doc.styles["GS Box Text"]
        K._run(p, "✦  " + label, 8.5, col, bold=True)
        if name.strip():
            p2 = c.add_paragraph(style="GS Box Text")
            K.add_runs(p2, name.strip(), 11, S.INK, bold=True)
            self._count(name)
            if desc.strip():
                K._run(p2, "  ·  " + desc.strip(), 9.5, S.SLATE, italic=True)
                self._count(desc)
        self._fill_cell(c, lines, col)
        K.no_split(t.rows[0])
        K.spacer(self.doc, 6)

    # ------------------------------------------------------------ tables
    def table(self, caption, lines, colour=None):
        col = colour or self.accent
        opts = {}
        rows = []
        for ln in lines:
            s = ln.strip()
            if not s:
                continue
            if s.startswith("@"):
                k, _, v = s[1:].partition(" ")
                opts[k] = v.strip()
                continue
            rows.append([x.strip() for x in s.split(" | ")])
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        if "widths" in opts:
            ws = [float(x) for x in opts["widths"].split(",")]
        else:
            ws = [1.0] * ncol
        tot = sum(ws)
        widths = [K.CONTENT_WIDTH_CM * w / tot for w in ws]
        header = opts.get("header", "on") != "off"
        size = float(opts.get("font", 9.5))
        if caption:
            self.table_no += 1
            cp = self.doc.add_paragraph(style="GS Table Caption")
            K.add_runs(cp, f"Table {self.table_no}  ·  {caption}")
            self._count(caption)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        K.table_indent_zero(t)
        t.style = self.doc.styles["Table Grid"]
        K.set_col_widths(t, widths)
        for i, r in enumerate(rows):
            for j, txt in enumerate(r):
                c = t.cell(i, j)
                K.cell_margins(c, 60, 60, 100, 100)
                K.cell_borders(c, top=(4, "D5DBE3"), bottom=(4, "D5DBE3"), left=(4, "D5DBE3"), right=(4, "D5DBE3"))
                p = c.paragraphs[0]
                p.style = self.doc.styles["GS Table Text"]
                parts = txt.split(" // ")
                for k, part in enumerate(parts):
                    pp = p if k == 0 else c.add_paragraph(style="GS Table Text")
                    if i == 0 and header:
                        K.add_runs(pp, part, size, "FFFFFF", bold=True, accent="FFFFFF")
                    else:
                        K.add_runs(pp, part, size, bold=(j == 0 and opts.get("firstbold") == "on"),
                                   accent=self.term_colour)
                    self._count(part)
                if i == 0 and header:
                    K.shade(c, col)
                elif i % 2 == 0:
                    K.shade(c, S.tint(col, 0.94))
            if i == 0 and header:
                K.repeat_header(t.rows[0])
            K.no_split(t.rows[i])
            if i == 0 or len(rows) <= 8:
                for cc in t.rows[i].cells:
                    for pp in cc.paragraphs:
                        pp.paragraph_format.keep_with_next = True
        K.spacer(self.doc, 6)

    # ------------------------------------------------------------ figures
    def figure_file(self, path, caption, width_in=None):
        if width_in is None:
            from PIL import Image
            with Image.open(path) as im:
                width_in = min(6.5, im.size[0] / 200 * 0.62)
        p = self.doc.add_paragraph(style="GS Figure")
        p.add_run().add_picture(str(path), width=Inches(width_in))
        self.fig_no += 1
        cp = self.doc.add_paragraph(style="Caption")
        K.add_runs(cp, f"Figure {self.fig_no}  ·  {caption}")
        self._count(caption)

    def figure(self, kind, caption, lines):
        path = F.render(kind, lines, self.code)
        width = None
        for ln in lines:
            if ln.strip().startswith("@width "):
                width = float(ln.split()[1])
        self.figure_file(path, caption, width)

    # ------------------------------------------------------------ contents
    def toc(self):
        self.doc.add_heading("Contents", 1)
        t, c = self._cell(S.tint(S.SLATE, 0.93), S.SLATE)
        p = c.paragraphs[0]
        p.style = self.doc.styles["GS Box Text"]
        K._run(p, "HOW TO BUILD THE CONTENTS", 9, S.SLATE, bold=True)
        for txt in ("The contents list below fills itself in. When Word asks \"This document contains fields that "
                    "may refer to other files. Update the fields?\", click Yes.",
                    "If it does not ask: right-click the grey contents area → Update Field → Update entire table → OK. "
                    "On a Mac, Control-click the same area. Or press Ctrl + A, then F9 (Windows), to update every "
                    "field at once. Repeat after any edit, before you export to PDF."):
            q = c.add_paragraph(style="GS Box Text")
            K.add_runs(q, txt)
        K.spacer(self.doc, 8)
        p = self.doc.add_paragraph()
        K.add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click here and choose Update Field to build the contents.")

    # ------------------------------------------------------------ the parser
    def run(self, lines):
        i, n = 0, len(lines)
        buf = []

        def flush():
            if buf:
                self.para(" ".join(x.strip() for x in buf))
                buf.clear()

        while i < n:
            ln = lines[i]
            s = ln.strip()
            if s.startswith("%%"):
                i += 1
                continue
            if not s:
                flush()
                i += 1
                continue
            if s.startswith(":::"):
                flush()
                head = s[3:].strip()
                j = i + 1
                body = []
                while j < n and lines[j].strip() != ":::":
                    body.append(lines[j])
                    j += 1
                self._block(head, body)
                i = j + 1
                continue
            m = re.match(r"^(#{1,5})\s+(.*)$", s)
            if m:
                flush()
                self.heading(len(m.group(1)), m.group(2).strip())
                i += 1
                continue
            if s == "@toc":
                flush()
                self.toc()
            elif s == "@pagebreak":
                flush()
                K.page_break(self.doc)
            elif s.startswith("@py "):
                flush()
                name, *args = s[4:].split()
                self.blocks[name](self, *args)
            elif re.match(r"^- ", ln):
                flush()
                self.bullet(s[2:], 1)
            elif re.match(r"^ {2,3}- ", ln):
                flush()
                self.bullet(s[2:], 2)
            elif re.match(r"^ {4,}- ", ln):
                flush()
                self.bullet(s[2:], 3)
            elif re.match(r"^(\d+|[ivx]+|[a-h])[.)]\s", s):
                flush()
                num, _, rest = s.partition(" ")
                self.numbered(num, rest.strip())
            elif s.startswith("> "):
                flush()
                self.para(s[2:], "GS Lead")
            else:
                buf.append(s)
            i += 1
        flush()

    def _block(self, head, body):
        kind, _, rest = head.partition(" ")
        kind = kind.strip()
        if "|" in kind:
            kind, rest = kind.split("|", 1)[0], "|" + kind.split("|", 1)[1] + rest
        rest = rest.strip()
        if kind in S.BOXES:
            title = rest.lstrip("|").strip()
            self.box(kind, title, body)
        elif kind in S.CARDS:
            self.card(kind, rest, body)
        elif kind == "table":
            cap = rest.lstrip("|").strip()
            self.table(cap, body)
        elif kind == "fig":
            ftype, _, cap = rest.partition("|")
            self.figure(ftype.strip(), cap.strip(), body)
        else:
            raise ValueError(f"Unknown block type: {kind!r}")
