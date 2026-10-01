"""The Prediction Papers: a replica of the FPSC CSS question paper, from content/PR/setN.txt.

Source file (blank lines and %% comments ignored):
    year: 2027
    ## MCQ
    stem || option a | option b | option c | option d || KEY || one-line reason
    ## PART-II
    Q2 || question text
    ...
    Q8 || Write short notes on any TWO of the following:
    Q8a || note text
    Q8b || note text
The paper is set in black Times New Roman, laid out as FPSC prints it: Roll Number box, the
four-line heading, the time-and-marks rules, the printed NOTE, then Part-I (Q. No. 1, twenty
MCQs, options (a)–(d)) and Part-II (Q. No. 2 to Q. No. 8, marks at the right margin).
"""
import re
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from engine import docxkit as K

FONT = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
W = K.CONTENT_WIDTH_CM
LETTERS = "abcd"


# ------------------------------------------------------------------ parsing
def load(path: Path):
    s = dict(year="2027", mcq=[], q={}, order=[])
    part = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        ln = raw.strip()
        if not ln or ln.startswith("%%"):
            continue
        if ln.startswith("year:"):
            s["year"] = ln.split(":", 1)[1].strip()
        elif ln == "## MCQ":
            part = "mcq"
        elif ln == "## PART-II":
            part = "q"
        elif part == "mcq":
            stem, opts, key, why = [x.strip() for x in ln.split("||")]
            opts = [o.strip() for o in opts.split("|")]
            if len(opts) != 4 or key.upper() not in "ABCD":
                raise SystemExit(f"{path.name}: bad MCQ → {stem[:60]}")
            s["mcq"].append(dict(stem=stem, opts=opts, key="ABCD".index(key.upper()), why=why))
        elif part == "q":
            k, txt = [x.strip() for x in ln.split("||", 1)]
            s["q"][k] = txt
            s["order"].append(k)
    if len(s["mcq"]) != 20:
        raise SystemExit(f"{path.name}: {len(s['mcq'])} MCQs, need 20")
    return s


# ------------------------------------------------------------------ low-level helpers
def _r(p, text, size=11, bold=False, italic=False, underline=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.color.rgb = BLACK
    r.bold, r.italic, r.underline = bold, italic, underline
    return r


def _runs(p, text, size=11, bold=False):
    """Plain text with *italic* spans (titles of books, Urdu terms)."""
    for i, part in enumerate(re.split(r"\*([^*]+)\*", text)):
        if part:
            _r(p, part, size, bold=bold, italic=bool(i % 2))


def _p(doc, align=None, before=0, after=0, left=0, hanging=0, keep=False, line=1.1):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    if left or hanging:
        f.left_indent = Cm(left)
        f.first_line_indent = Cm(-hanging)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    f.keep_with_next = keep
    return p


def _right_tab(p, cm=W):
    p.paragraph_format.tab_stops.add_tab_stop(Cm(cm), WD_TAB_ALIGNMENT.RIGHT)


def _rule(doc, sz=12, after=2):
    p = _p(doc, after=after)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", "1"), ("w:color", "000000")):
        e.set(qn(k), v)
    b.append(e)
    pPr.append(b)
    return p


def _plain_table(doc, rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols)
    K.table_indent_zero(t)
    K.set_col_widths(t, widths)
    for row in t.rows:
        for c in row.cells:
            K.cell_margins(c, 10, 10, 40, 40)
    return t


def _cell_text(c, text, size=11, bold=False, align=None):
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, text, size, bold)
    return p


# ------------------------------------------------------------------ the paper
def _roll_box(doc):
    t = _plain_table(doc, 1, 3, [W - 6.4, 2.4, 4.0])
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    _cell_text(t.cell(0, 1), "Roll Number", 10, True, "r")
    box = t.cell(0, 2)
    edge = (8, "000000")
    K.cell_borders(box, top=edge, left=edge, bottom=edge, right=edge)
    _cell_text(box, " ", 12)


def _heading(doc, year):
    _roll_box(doc)
    for txt, size, after in (("FEDERAL PUBLIC SERVICE COMMISSION", 14, 0),
                             (f"COMPETITIVE EXAMINATION-{year}", 12, 0),
                             ("FOR RECRUITMENT TO POSTS IN BS-17", 12, 0),
                             ("UNDER THE FEDERAL GOVERNMENT", 12, 4)):
        _r(_p(doc, "c", after=after), txt, size, bold=True)
    _r(_p(doc, "c", after=4), "GENDER STUDIES", 13, bold=True, underline=True)
    _rule(doc, 12, 0)
    for left, right in (("TIME ALLOWED: THREE HOURS", "TOTAL MARKS: 100"),
                        ("PART-I (MCQS): MAXIMUM 30 MINUTES", "PART-I (MCQS): MAXIMUM MARKS = 20"),
                        ("PART-II", "PART-II: MAXIMUM MARKS = 80")):
        p = _p(doc, before=2)
        _right_tab(p)
        _r(p, left, 10.5, bold=True)
        _r(p, "\t" + right, 10.5, bold=True)
    _rule(doc, 12, 6)


def _note(doc, items):
    t = _plain_table(doc, len(items), 2, [1.4, W - 1.4])
    for i, txt in enumerate(items):
        if i == 0:
            _cell_text(t.cell(i, 0), "NOTE:", 10.5, True)
        p = _cell_text(t.cell(i, 1), "", 11)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
        _r(p, f"({'i ii iii iv v vi vii'.split()[i]})\t", 10.5)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.8))
        for j, part in enumerate(re.split(r"\^\^(.+?)\^\^", txt)):
            if part:
                _r(p, part, 10.5, bold=bool(j % 2))
    _p(doc, after=2)


def _stars(doc):
    _r(_p(doc, "c", before=6), "**********", 12, bold=True)


def _options_table(doc, opts):
    longest = max(len(o) for o in opts)
    cols = 4 if longest <= 22 else 2 if longest <= 48 else 1
    rows = 4 // cols
    t = _plain_table(doc, rows, cols, [(W - 0.9) / cols] * cols)
    # indent the grid under the stem
    tblPr = t._tbl.tblPr
    ind = tblPr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tblPr.append(ind)
    ind.set(qn("w:w"), str(int(0.9 / 2.54 * 1440)))
    ind.set(qn("w:type"), "dxa")
    for i, o in enumerate(opts):
        c = t.cell(i // cols, i % cols)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.65)
        _r(p, f"({LETTERS[i]}) ", 10.5)
        _runs(p, o, 10.5)
    for ri, row in enumerate(t.rows):
        K.no_split(row)
        if ri < len(t.rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    return t


def render_part1(b, s):
    doc = b.doc
    K.page_break(doc)
    _heading(doc, s["year"])
    _note(doc, ["First attempt ^^PART-I (MCQS)^^ on separate Answer Sheet which shall be taken back after 30 minutes.",
                "Overwriting/cutting of the options/answers will not be given credit."])
    _r(_p(doc, "c", after=6), "PART-I (MCQS) (COMPULSORY)", 12, bold=True, underline=True)
    p = _p(doc, left=1.9, hanging=1.9, after=0)
    _right_tab(p)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.9))
    _r(p, "Q. No. 1.\t", 11, bold=True)
    _r(p, "(i) Select the best option/answer and fill in the appropriate box on the Answer Sheet.", 11)
    p = _p(doc, left=1.9, after=8)
    _right_tab(p)
    _r(p, "(ii) Answers given anywhere else, other than Answer Sheet, will not be considered.", 11)
    _r(p, "\t(20×1=20)", 11, bold=True)
    for i, m in enumerate(s["mcq"], 1):
        p = _p(doc, left=0.9, hanging=0.9, before=3, keep=True)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
        _r(p, f"{i}.\t", 11)
        _runs(p, m["stem"], 11)
        _options_table(doc, m["opts"])
    _stars(doc)


def render_part2(b, s):
    doc = b.doc
    K.page_break(doc)
    _heading(doc, s["year"])
    _note(doc, ["Part-II is to be attempted on the separate Answer Book.",
                "Attempt ^^ONLY FOUR^^ questions from ^^PART-II^^. ^^ALL^^ questions carry ^^EQUAL^^ marks.",
                "All the parts (if any) of each Question must be attempted at one place instead of at different places.",
                "Write Q. No. in the Answer Book in accordance with Q. No. in the Q. Paper.",
                "No Page/Space be left blank between the answers. All the blank pages of Answer Book must be crossed.",
                "Extra attempt of any question or any part of the question will not be considered."])
    _r(_p(doc, "c", after=4), "PART-II", 12, bold=True, underline=True)
    last = None
    for k in s["order"]:
        txt = s["q"][k]
        m = re.match(r"Q(\d)([a-c]?)$", k)
        n, part = m.group(1), m.group(2)
        if not part:
            notes = any(o.startswith(f"Q{n}") and len(o) == 3 for o in s["order"])
            p = _p(doc, "j", left=1.9, hanging=1.9, before=4, after=2, keep=notes, line=1.05)
            _right_tab(p)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.9))
            _r(p, f"Q. No. {n}.\t", 11, bold=True)
            _runs(p, txt, 11)
            _r(p, "\t(10+10)" if notes else "\t(20)", 11, bold=True)
        else:
            p = _p(doc, "j", left=2.8, hanging=0.9, after=3, line=1.05)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(2.8))
            _r(p, f"({part})\t", 11)
            _runs(p, txt, 11)
        last = p
    last.paragraph_format.keep_with_next = True
    _stars(doc)


def question_box(b, s, key, set_no):
    """The predicted question as a box at the head of its model answer."""
    txt = s["q"][key]
    m = re.match(r"Q(\d)([a-c]?)$", key)
    n, part = m.group(1), m.group(2)
    if part:
        title = f"Prediction Paper Set {set_no} · Q. No. {n}({part}) · short note · 10 marks"
        txt = f"{s['q'][f'Q{n}'].rstrip(':')}: {txt}"
    else:
        title = f"Prediction Paper Set {set_no} · Q. No. {n} · 20 marks"
    b.box("glance", title, [txt])
