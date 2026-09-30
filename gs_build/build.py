#!/usr/bin/env python3
"""Build a Gender Studies Word document from its GSM source.

    python3 gs_build/build.py MA            # the Master Anatomy
    python3 gs_build/build.py T1KB T1QA     # a topic's two documents
    python3 gs_build/build.py all           # every document whose source exists

Sources live in gs_build/content/<CODE-or-topic>/ ; outputs go into "CSS Gender Studies/".
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx.enum.table import WD_ROW_HEIGHT_RULE          # noqa: E402
from docx.shared import Cm, Pt                          # noqa: E402

from engine import docxkit as K                          # noqa: E402
from engine import style as S                            # noqa: E402
from engine.markup import Builder, read_source           # noqa: E402
from data import topics as TP                            # noqa: E402
import blocks                                            # noqa: E402

CONTENT = Path(__file__).resolve().parent / "content"
UPDATED = "September 2026"


def registry():
    reg = {
        "MA": dict(src=CONTENT / "MA" / "main.gsm", folder="MA_Master Anatomy",
                   file="MA_CSS_Gender_Studies_Master_Anatomy.docx", colour=S.PLUM),
        "QA": dict(src=CONTENT / "QA" / "main.gsm", folder="QA_The Question Answers",
                   file="QA_The Question Answers — All Seven Topics.docx", colour=S.PLUM),
        "OL": dict(src=CONTENT / "OL" / "main.gsm", folder="OL_The One-Liner and MCQ Bank",
                   file="OL_The One-Liner and MCQ Bank.docx", colour=S.PLUM),
        "RN": dict(src=CONTENT / "RN" / "main.gsm", folder="RN_The Revision Notes",
                   file="RN_The Revision Notes — All Seven Topics.docx", colour=S.PLUM),
        "PR": dict(src=CONTENT / "PR" / "main.gsm", folder="PR_The Prediction Papers",
                   file="PR_The Prediction Papers — Sets 1–3.docx", colour=S.PLUM),
        "FB": dict(src=CONTENT / "FB" / "main.gsm", folder="FB_The Fact Book",
                   file="FB_The Fact Book.docx", colour=S.PLUM),
    }
    reg["TEST"] = dict(src=CONTENT / "TEST" / "main.gsm", folder=str(S.CACHE), file="TEST.docx", colour=S.PLUM)
    for t in TP.TOPICS:
        n = t["n"]
        reg[f"T{n}KB"] = dict(src=CONTENT / f"T{n}" / "KB.gsm", folder=TP.folder(n), file=TP.kb_file(n),
                              colour=t["colour"])
        reg[f"T{n}QA"] = dict(src=CONTENT / f"T{n}" / "QA.gsm", folder=TP.folder(n), file=TP.qa_file(n),
                              colour=t["colour"])
    return reg


def cover(doc, fm, accent):
    t = doc.add_table(rows=3, cols=1)
    K.table_indent_zero(t)
    K.set_col_widths(t, [K.CONTENT_WIDTH_CM])
    top, mid, bot = (t.cell(i, 0) for i in range(3))
    for c in (top, mid, bot):
        K.cell_borders(c)
    K.shade(top, accent)
    K.cell_margins(top, 500, 400, 500, 500)
    t.rows[0].height = Cm(13.5)
    t.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    p = top.paragraphs[0]
    K._run(p, "CSS  GENDER  STUDIES", 11, S.tint(accent, 0.55), bold=True)
    p.paragraph_format.space_after = Pt(40)
    p = top.add_paragraph()
    K._run(p, fm.get("code", ""), 54, "FFFFFF", bold=True)
    p.paragraph_format.space_after = Pt(0)
    p = top.add_paragraph()
    K._run(p, fm.get("title", ""), 30, "FFFFFF", bold=True)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(14)
    if fm.get("subtitle"):
        p = top.add_paragraph()
        K._run(p, fm["subtitle"], 13.5, S.tint(accent, 0.75))
        p.paragraph_format.line_spacing = 1.15
    K.shade(mid, S.tint(accent, 0.93))
    K.cell_margins(mid, 300, 300, 500, 500)
    p = mid.paragraphs[0]
    K._run(p, "WHAT THIS DOCUMENT GIVES YOU", 9.5, accent, bold=True)
    p.paragraph_format.space_after = Pt(6)
    for item in [x.strip() for x in fm.get("gives", "").split(";") if x.strip()]:
        q = mid.add_paragraph()
        K._run(q, "◆  ", 10, accent, bold=True)
        K.add_runs(q, item, 10.5)
        q.paragraph_format.space_after = Pt(3)
    K.cell_margins(bot, 240, 100, 500, 500)
    p = bot.paragraphs[0]
    K._run(p, "Federal Public Service Commission  ·  CSS Competitive Examination  ·  Optional paper, 100 marks",
           9, S.SLATE)
    p = bot.add_paragraph()
    K._run(p, f"Written for CE-2027 onwards  ·  Record: CE-2016 to CE-2026  ·  Facts current to {UPDATED}",
           9, S.SLATE, italic=True)


def build(code):
    reg = registry()
    if code not in reg:
        raise SystemExit(f"Unknown document code {code}. Known: {', '.join(reg)}")
    spec = reg[code]
    fm, lines = read_source(spec["src"])
    accent = spec["colour"]
    doc = K.new_document(accent)
    K.set_properties(doc, fm.get("title", code), fm.get("subtitle", ""))
    K.header_footer(doc, f"CSS Gender Studies  ·  {fm.get('header', fm.get('title', ''))}", fm.get("code", code),
                    accent)
    cover(doc, fm, accent)
    b = Builder(doc, code, accent, blocks.REGISTRY)
    b.run(lines)
    out_dir = S.SUBJECT_DIR / spec["folder"]
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / spec["file"]
    doc.save(out)
    print(f"{code}: {b.words:,} words · {b.fig_no} figures · {b.table_no} captioned tables -> {out.relative_to(S.REPO)}")
    return out


if __name__ == "__main__":
    codes = sys.argv[1:] or ["MA"]
    if codes == ["all"]:
        codes = [c for c, s in registry().items() if s["src"].exists()]
    for c in codes:
        build(c)
